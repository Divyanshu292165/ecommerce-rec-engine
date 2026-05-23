from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    
    # 16:9 Aspect Ratio
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Colors
    bg_color = RGBColor(15, 23, 42) # Slate 900
    text_color = RGBColor(248, 250, 252) # Slate 50
    accent_color = RGBColor(168, 85, 247) # Purple 500
    sub_color = RGBColor(148, 163, 184) # Slate 400

    def apply_dark_theme(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = bg_color

    def add_title_slide():
        slide_layout = prs.slide_layouts[6] # Blank
        slide = prs.slides.add_slide(slide_layout)
        apply_dark_theme(slide)
        
        # Title
        txBox = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11.333), Inches(1.5))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = "NeuralShop RecEngine"
        p.font.size = Pt(64)
        p.font.bold = True
        p.font.color.rgb = text_color
        p.alignment = PP_ALIGN.CENTER
        
        # Subtitle
        p2 = tf.add_paragraph()
        p2.text = "AI-Powered E-Commerce Recommendation & Deal Engine"
        p2.font.size = Pt(32)
        p2.font.color.rgb = accent_color
        p2.alignment = PP_ALIGN.CENTER

        # Footer
        txBox2 = slide.shapes.add_textbox(Inches(1), Inches(6.5), Inches(11.333), Inches(0.5))
        tf2 = txBox2.text_frame
        p3 = tf2.paragraphs[0]
        p3.text = "Real-Time Amazon Data \u2022 Deal Analysis \u2022 Glassmorphism UI"
        p3.font.size = Pt(18)
        p3.font.color.rgb = sub_color
        p3.alignment = PP_ALIGN.CENTER

    def add_content_slide(title_text, bullets):
        slide_layout = prs.slide_layouts[6] # Blank
        slide = prs.slides.add_slide(slide_layout)
        apply_dark_theme(slide)
        
        # Title
        txBox = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11.333), Inches(1))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(44)
        p.font.bold = True
        p.font.color.rgb = accent_color
        
        # Line separator
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(1), Inches(1.6), Inches(11.333), Inches(0.02)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = accent_color
        shape.line.fill.background()

        # Content
        txBox2 = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(11.333), Inches(5))
        tf2 = txBox2.text_frame
        tf2.word_wrap = True
        
        for i, bullet in enumerate(bullets):
            p = tf2.add_paragraph()
            p.text = "\u2022 " + bullet
            p.font.size = Pt(28)
            p.font.color.rgb = text_color
            p.space_after = Pt(20)

    # Slide 1: Title
    add_title_slide()

    # Slide 2: The Problem
    add_content_slide("The Problem", [
        "E-commerce platforms overwhelm users with thousands of choices.",
        "Prices fluctuate constantly, making it hard to spot genuine discounts.",
        "Users waste hours checking multiple tabs to verify 'deal' authenticity.",
        "Static homepages show irrelevant, outdated recommendations."
    ])

    # Slide 3: The Solution
    add_content_slide("The Solution: NeuralShop", [
        "A real-time recommendation engine connected directly to Amazon India.",
        "Fetches live pricing, ratings, and stock status instantly.",
        "Smart AI Deal Analysis algorithms categorize items (Buy Now, Wait).",
        "A beautiful, responsive Glassmorphism UI that prioritizes user experience."
    ])

    # Slide 4: Key Features
    add_content_slide("Key Features", [
        "Live INR Pricing: Automatically converts and formats Amazon prices.",
        "Deal Analysis Badges: Mathematically calculates discounts against MSRP.",
        "Trending Recommendations: Auto-loads 'best laptops', 'smartphones', etc.",
        "Lightning-Fast Search: RapidAPI integration for instant query results."
    ])

    # Slide 5: Tech Stack Architecture
    add_content_slide("Technology Stack", [
        "Frontend: Vanilla JavaScript, HTML5, CSS3 with advanced backdrop-filters.",
        "Backend: FastAPI (Python) for asynchronous, high-performance endpoints.",
        "Data Source: Real-Time Amazon Data via RapidAPI (Live scraping proxy).",
        "Deployment: Render Cloud Hosting with automatic CI/CD from GitHub."
    ])

    # Slide 6: Future Scope
    add_content_slide("Future Scope", [
        "Price Drop Alerts: Email notifications when a 'Wait' item drops in price.",
        "User Profiles: Save favorite searches and track products over time.",
        "Multi-Store Comparison: Compare Amazon, Flipkart, and Croma in one click.",
        "Generative AI Chatbot: 'Ask NeuralShop what phone to buy under \u20b920,000'."
    ])

    # Slide 7: Thank You
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    apply_dark_theme(slide)
    txBox = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(11.333), Inches(1.5))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "Thank You!"
    p.font.size = Pt(64)
    p.font.bold = True
    p.font.color.rgb = accent_color
    p.alignment = PP_ALIGN.CENTER
    
    p2 = tf.add_paragraph()
    p2.text = "https://ecommerce-rec-engine.onrender.com"
    p2.font.size = Pt(24)
    p2.font.color.rgb = sub_color
    p2.alignment = PP_ALIGN.CENTER

    prs.save("c:/Users/Divyanshu/OneDrive/Desktop/NeuralShop_Presentation.pptx")
    print("Presentation saved!")

if __name__ == "__main__":
    create_presentation()
